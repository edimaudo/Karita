from __future__ import annotations

import csv
import hashlib
import io
import re
from pathlib import Path
from typing import Any

from docx import Document
from fastapi import UploadFile
from openpyxl import load_workbook
from pypdf import PdfReader


ALLOWED_EXTENSIONS = {".pdf", ".docx", ".csv", ".xlsx"}
ALLOWED_CONTENT_TYPES = {
    ".pdf": {"application/pdf", "application/octet-stream"},
    ".docx": {
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application/octet-stream",
    },
    ".csv": {"text/csv", "text/plain", "application/csv", "application/octet-stream"},
    ".xlsx": {
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "application/octet-stream",
    },
}

MAX_FILE_SIZE = 10 * 1024 * 1024
MAX_TEXT_CHARS_PER_FILE = 100_000

# This is intentionally a warning signal, not a security boundary.
PROMPT_INJECTION_PATTERNS = (
    re.compile(r"ignore\s+(all\s+)?previous\s+instructions", re.I),
    re.compile(r"disregard\s+(all\s+)?previous\s+instructions", re.I),
    re.compile(r"reveal\s+(the\s+)?system\s+prompt", re.I),
    re.compile(r"show\s+(me\s+)?your\s+system\s+instructions", re.I),
    re.compile(r"you\s+are\s+now\s+(a|an)\s+", re.I),
)


def sanitize_filename(filename: str | None) -> str:
    """Return a safe display/storage filename without path components."""
    name = Path(filename or "upload").name
    name = re.sub(r"[^A-Za-z0-9._-]", "_", name)
    name = name.strip(".") or "upload"
    return name[:180]


def file_extension(filename: str | None) -> str:
    return Path(filename or "").suffix.lower()


def validate_upload(upload: UploadFile, content: bytes) -> None:
    """Validate extension, declared MIME type and size.

    Extension and MIME checks are deliberately separate. A production deployment
    should add a malware scanner and stronger content-signature validation.
    """
    extension = file_extension(upload.filename)
    if extension not in ALLOWED_EXTENSIONS:
        raise ValueError(f"Unsupported file type: {sanitize_filename(upload.filename)}")

    if len(content) > MAX_FILE_SIZE:
        raise ValueError(
            f"File exceeds the {MAX_FILE_SIZE // (1024 * 1024)} MB limit: "
            f"{sanitize_filename(upload.filename)}"
        )

    content_type = (upload.content_type or "application/octet-stream").split(";")[0].strip().lower()
    if content_type not in ALLOWED_CONTENT_TYPES[extension]:
        raise ValueError(
            f"File type does not match its declared content type: "
            f"{sanitize_filename(upload.filename)}"
        )


def _decode_text(data: bytes) -> str:
    """Decode common text encodings without failing the whole engagement."""
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _extract_pdf(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    pages = []
    for number, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        pages.append(f"[Page {number}]\n{text.strip()}")
    return "\n\n".join(pages)


def _extract_docx(data: bytes) -> str:
    document = Document(io.BytesIO(data))
    sections: list[str] = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            sections.append(text)

    for table_number, table in enumerate(document.tables, start=1):
        sections.append(f"[Table {table_number}]")
        for row in table.rows:
            sections.append(" | ".join(cell.text.strip() for cell in row.cells))

    return "\n".join(sections)


def _extract_csv(data: bytes) -> str:
    text = _decode_text(data)
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample)
    except csv.Error:
        dialect = csv.excel

    rows = list(csv.reader(io.StringIO(text), dialect))
    if not rows:
        return ""

    # Markdown-like tabular representation is compact and easy for an LLM to use.
    output: list[str] = []
    headers = [cell.strip() for cell in rows[0]]
    output.append(" | ".join(headers))
    output.append(" | ".join("---" for _ in headers))
    for row in rows[1:]:
        values = [cell.strip() for cell in row]
        if len(values) < len(headers):
            values.extend([""] * (len(headers) - len(values)))
        output.append(" | ".join(values[: len(headers)]))
    return "\n".join(output)


def _extract_xlsx(data: bytes) -> str:
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sections: list[str] = []

    try:
        for worksheet in workbook.worksheets:
            sections.append(f"[Worksheet: {worksheet.title}]")
            rows = worksheet.iter_rows(values_only=True)
            try:
                first_row = next(rows)
            except StopIteration:
                continue

            headers = ["" if value is None else str(value).strip() for value in first_row]
            sections.append(" | ".join(headers))
            sections.append(" | ".join("---" for _ in headers))

            for row in rows:
                values = ["" if value is None else str(value).strip() for value in row]
                if not any(values):
                    continue
                if len(values) < len(headers):
                    values.extend([""] * (len(headers) - len(values)))
                sections.append(" | ".join(values[: len(headers)]))
            sections.append("")
    finally:
        workbook.close()

    return "\n".join(sections)


def extract_text(filename: str, data: bytes) -> str:
    """Extract usable text/data from a supported Karita file."""
    extension = file_extension(filename)

    if extension == ".pdf":
        text = _extract_pdf(data)
    elif extension == ".docx":
        text = _extract_docx(data)
    elif extension == ".csv":
        text = _extract_csv(data)
    elif extension == ".xlsx":
        text = _extract_xlsx(data)
    else:
        raise ValueError(f"Unsupported file type: {extension or 'unknown'}")

    text = text.strip()
    return text[:MAX_TEXT_CHARS_PER_FILE]


def detect_prompt_injection(text: str) -> list[str]:
    """Return suspicious instruction-like phrases found in extracted content."""
    findings: list[str] = []
    for pattern in PROMPT_INJECTION_PATTERNS:
        match = pattern.search(text)
        if match:
            findings.append(match.group(0)[:200])
    return findings


def build_evidence(filename: str, data: bytes, content_type: str | None = None) -> dict[str, Any]:
    """Create a normalized evidence object for an uploaded document."""
    safe_name = sanitize_filename(filename)
    extension = file_extension(safe_name)
    text = extract_text(safe_name, data)
    injection_flags = detect_prompt_injection(text)

    return {
        "filename": safe_name,
        "extension": extension,
        "content_type": content_type or "application/octet-stream",
        "size_bytes": len(data),
        "sha256": hashlib.sha256(data).hexdigest(),
        "text": text,
        "prompt_injection_detected": bool(injection_flags),
        "prompt_injection_findings": injection_flags,
    }


async def read_upload(upload: UploadFile) -> dict[str, Any]:
    """Read, validate and normalize one FastAPI UploadFile."""
    data = await upload.read(MAX_FILE_SIZE + 1)
    validate_upload(upload, data)
    return build_evidence(upload.filename or "upload", data, upload.content_type)
